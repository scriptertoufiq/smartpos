"""M11 tests: reports, profit and exports.

Acceptance criteria from docs/PLAN.md:
  * every report's totals match a hand-written SQL check on the same data
  * profit is unchanged after a product's purchase price is edited
  * reports exclude held and cancelled sales
  * a 12-month report over 10,000 sales renders in under 3 seconds
  * exported PDF and Excel contain the same rows as the on-screen table
"""

from __future__ import annotations

import time
from datetime import date, datetime, timedelta
from decimal import Decimal

import pytest

from smartpos.core.exceptions import DatabaseError
from smartpos.db import connection
from smartpos.reports import exporters
from smartpos.services.auth_service import AuthService
from smartpos.services.customer_service import CustomerService
from smartpos.services.product_service import ProductService
from smartpos.services.purchase_service import PurchaseDraft, PurchaseLine, PurchaseService
from smartpos.services.report_service import ReportService
from smartpos.services.sale_return_service import ReturnDraft, ReturnLine, SaleReturnService
from smartpos.services.sale_service import SaleDraft, SaleLine, SaleService
from smartpos.services.supplier_service import SupplierService
from smartpos.services.warehouse_service import WarehouseService

MARK = "ZZRPT"
TODAY = date.today()


@pytest.fixture(scope="module")
def db():
    try:
        connection.init_pool()
        connection.check_connection()
    except DatabaseError as exc:
        pytest.skip(f"MySQL not reachable: {exc}")
    yield
    connection.close_pool()


def _purge() -> None:
    mine = ("SELECT * FROM (SELECT DISTINCT i.sale_id FROM sale_items i "
            "JOIN products p ON p.id = i.product_id WHERE p.sku LIKE %s) AS t")
    with connection.transaction() as conn:
        cur = conn.cursor()
        # Order matters: a payment row has no foreign key to the document it
        # names, so deleting the document first leaves the payment dangling --
        # and a dangling payment silently skews "collected today" and every due.
        my_returns = (f"SELECT * FROM (SELECT id FROM sale_returns "
                      f"WHERE sale_id IN ({mine})) AS r")
        cur.execute(f"DELETE FROM customer_payments WHERE reference_type = 'sale_return' "
                    f"AND reference_id IN ({my_returns})", (MARK + "%",))
        cur.execute(f"DELETE FROM sale_return_items WHERE sale_return_id IN "
                    f"({my_returns})", (MARK + "%",))
        cur.execute(f"DELETE FROM sale_returns WHERE sale_id IN ({mine})", (MARK + "%",))
        # A counter sale belongs to the walk-in customer, so the customer filter
        # further down never reaches its payment.
        cur.execute(f"DELETE FROM customer_payments WHERE reference_type = 'sale' "
                    f"AND reference_id IN ({mine})", (MARK + "%",))
        cur.execute(f"DELETE FROM sales WHERE id IN ({mine})", (MARK + "%",))
        for sql in (
            "DELETE FROM customer_payments WHERE customer_id IN "
            " (SELECT id FROM customers WHERE name LIKE %s)",
            "DELETE FROM sale_items WHERE product_id IN "
            " (SELECT id FROM products WHERE sku LIKE %s)",
            "DELETE FROM purchase_items WHERE product_id IN "
            " (SELECT id FROM products WHERE sku LIKE %s)",
            "DELETE FROM purchases WHERE supplier_id IN "
            " (SELECT id FROM suppliers WHERE name LIKE %s)",
            "DELETE FROM supplier_payments WHERE supplier_id IN "
            " (SELECT id FROM suppliers WHERE name LIKE %s)",
            "DELETE FROM stock_movements WHERE product_id IN "
            " (SELECT id FROM products WHERE sku LIKE %s)",
            "DELETE FROM stock WHERE product_id IN "
            " (SELECT id FROM products WHERE sku LIKE %s)",
            "DELETE FROM products WHERE sku LIKE %s",
            "DELETE FROM suppliers WHERE name LIKE %s",
            "DELETE FROM customers WHERE name LIKE %s AND is_walk_in = 0",
        ):
            cur.execute(sql, (MARK + "%",))
        cur.close()


@pytest.fixture
def ctx(db):
    _purge()
    auth = AuthService()
    auth.login("admin", "admin123")
    products = ProductService(auth)
    sales = SaleService(auth)

    warehouse = next(w for w in WarehouseService(auth).list_all()
                     if w.code == "MAIN").id
    supplier = SupplierService(auth).create(f"{MARK} Supplier")
    customer = CustomerService(auth).create(f"{MARK} Rahim")
    rice = products.create(f"{MARK}-RICE", "Rice 5kg", sale_price="420.00")
    oil = products.create(f"{MARK}-OIL", "Oil 1L", sale_price="175.00")

    PurchaseService(auth).create(PurchaseDraft(
        supplier_id=supplier, warehouse_id=warehouse,
        lines=[PurchaseLine(rice, Decimal("200"), Decimal("300.00")),
               PurchaseLine(oil, Decimal("200"), Decimal("100.00"))],
        paid_amount=Decimal("10000.00")))

    yield {"auth": auth, "reports": ReportService(auth), "sales": sales,
           "returns": SaleReturnService(auth), "products": products,
           "customers": CustomerService(auth),
           "warehouse": warehouse, "rice": rice, "oil": oil,
           "supplier": supplier, "customer": customer}
    _purge()


def _sell(ctx, rice=2, oil=0, tendered=None, customer_id=None, discount="0"):
    lines = [SaleLine(ctx["rice"], Decimal(str(rice)), Decimal("420.00"))]
    if oil:
        lines.append(SaleLine(ctx["oil"], Decimal(str(oil)), Decimal("175.00")))
    gross = sum(l.quantity * l.unit_price for l in lines) - Decimal(discount)
    return ctx["sales"].create(SaleDraft(
        warehouse_id=ctx["warehouse"], customer_id=customer_id, lines=lines,
        discount_value=Decimal(discount),
        tendered=Decimal(tendered) if tendered is not None else gross))


def _sql(query: str, params=()) -> Decimal:
    """A hand-written check, deliberately not sharing the service's code."""
    with connection.get_connection() as conn:
        cur = conn.cursor()
        try:
            cur.execute(query, params)
            row = cur.fetchone()
            return Decimal(str(row[0] or 0))
        finally:
            cur.close()


# --------------------------------------------------- totals vs hand-written SQL

def test_sales_totals_match_a_hand_written_query(ctx):
    """Acceptance criterion."""
    _sell(ctx, rice=3)
    _sell(ctx, rice=1, oil=2)
    _sell(ctx, rice=5, tendered="0")              # a credit sale

    report = ctx["reports"].sales_report(TODAY, TODAY, ctx["warehouse"])
    expected_billed = _sql(
        "SELECT COALESCE(SUM(grand_total),0) FROM sales "
        "WHERE status='completed' AND DATE(sale_date)=%s AND warehouse_id=%s",
        (TODAY, ctx["warehouse"]))
    expected_paid = _sql(
        "SELECT COALESCE(SUM(paid_amount),0) FROM sales "
        "WHERE status='completed' AND DATE(sale_date)=%s AND warehouse_id=%s",
        (TODAY, ctx["warehouse"]))

    assert report.totals["billed"] == expected_billed
    assert report.totals["collected"] == expected_paid
    assert report.totals["outstanding"] == expected_billed - expected_paid
    assert report.totals["invoices"] == len(report.rows)


def test_stock_report_value_matches_a_hand_written_query(ctx):
    report = ctx["reports"].stock_report(ctx["warehouse"])
    expected = _sql(
        "SELECT COALESCE(SUM(st.quantity * p.purchase_price),0) FROM stock st "
        "JOIN products p ON p.id=st.product_id "
        "WHERE p.is_active=1 AND st.warehouse_id=%s", (ctx["warehouse"],))
    assert report.totals["value"] == expected


def test_customer_dues_match_the_customer_service(ctx):
    """Two independent implementations of the same figure must agree."""
    cid = ctx["customer"]
    _sell(ctx, rice=4, tendered="0", customer_id=cid)
    ctx["customers"].collect_due(cid, "500.00")

    report = ctx["reports"].customer_dues()
    row = next(r for r in report.rows if r["customer"] == f"{MARK} Rahim")
    assert Decimal(str(row["due"])) == ctx["customers"].due_balance(cid)


def test_supplier_payables_match_the_supplier_service(ctx):
    from smartpos.services.supplier_service import SupplierService
    suppliers = SupplierService(ctx["auth"])
    report = ctx["reports"].supplier_payables()
    row = next(r for r in report.rows if r["supplier"] == f"{MARK} Supplier")
    assert Decimal(str(row["payable"])) == suppliers.payable_balance(ctx["supplier"])


def test_purchase_report_totals_match(ctx):
    report = ctx["reports"].purchase_report(TODAY, TODAY)
    mine = [r for r in report.rows if r["supplier"] == f"{MARK} Supplier"]
    assert len(mine) == 1
    assert Decimal(str(mine[0]["grand_total"])) == Decimal("80000.00")
    assert Decimal(str(mine[0]["paid_amount"])) == Decimal("10000.00")
    assert mine[0]["payment_status"] == "partial"


# ------------------------------------------------------------------- profit

def test_profit_is_revenue_less_the_cost_stamped_on_each_line(ctx):
    _sell(ctx, rice=10)                      # sells at 420, cost stamped at 300
    figures = ctx["reports"].profit_loss(TODAY, TODAY, ctx["warehouse"])
    assert figures["line_revenue"] == Decimal("4200.00")
    assert figures["cogs"] == Decimal("3000.00")
    assert figures["gross_profit"] == Decimal("1200.00")


def test_profit_is_unchanged_when_the_purchase_price_is_edited(ctx):
    """Acceptance criterion -- the whole reason unit_cost is stored."""
    _sell(ctx, rice=10)
    before = ctx["reports"].profit_loss(TODAY, TODAY, ctx["warehouse"])

    ctx["products"].update(ctx["rice"], purchase_price="999.00")

    after = ctx["reports"].profit_loss(TODAY, TODAY, ctx["warehouse"])
    assert after["cogs"] == before["cogs"]
    assert after["gross_profit"] == before["gross_profit"] == Decimal("1200.00")


def test_a_return_removes_both_the_revenue_and_the_cost(ctx):
    sale_id = _sell(ctx, rice=10)
    item_id = ctx["returns"].returnable_quantities(sale_id)[0].sale_item_id
    ctx["returns"].create(ReturnDraft(sale_id, [ReturnLine(item_id, Decimal("4"))]))

    figures = ctx["reports"].profit_loss(TODAY, TODAY, ctx["warehouse"])
    assert figures["returned"] == Decimal("1680.00")       # 4 x 420
    assert figures["returned_cogs"] == Decimal("1200.00")  # 4 x 300
    assert figures["net_revenue"] == Decimal("2520.00")
    assert figures["net_cogs"] == Decimal("1800.00")
    assert figures["gross_profit"] == Decimal("720.00")    # 6 units x 120 margin


def test_an_invoice_discount_comes_off_the_profit(ctx):
    _sell(ctx, rice=10, discount="200")
    figures = ctx["reports"].profit_loss(TODAY, TODAY, ctx["warehouse"])
    assert figures["header_discount"] == Decimal("200.00")
    assert figures["gross_profit"] == Decimal("1000.00")   # 1200 - 200


def test_tax_is_reported_but_is_not_profit(ctx):
    draft = SaleDraft(warehouse_id=ctx["warehouse"],
                      lines=[SaleLine(ctx["rice"], Decimal("10"), Decimal("420.00"))],
                      tax_amount=Decimal("630.00"), tendered=Decimal("4830.00"))
    ctx["sales"].create(draft)
    figures = ctx["reports"].profit_loss(TODAY, TODAY, ctx["warehouse"])
    assert figures["tax_collected"] == Decimal("630.00")
    assert figures["gross_profit"] == Decimal("1200.00"), "tax must not inflate profit"


# ------------------------------------------------- held and cancelled sales

def test_reports_exclude_held_sales(ctx):
    """Acceptance criterion."""
    _sell(ctx, rice=2)
    ctx["sales"].hold(SaleDraft(
        warehouse_id=ctx["warehouse"],
        lines=[SaleLine(ctx["rice"], Decimal("50"), Decimal("420.00"))]))

    report = ctx["reports"].sales_report(TODAY, TODAY, ctx["warehouse"])
    assert report.totals["invoices"] == 1
    assert report.totals["billed"] == Decimal("840.00")


def test_reports_exclude_cancelled_sales(ctx):
    sale_id = _sell(ctx, rice=2)
    _sell(ctx, rice=3)
    with connection.transaction() as conn:
        cur = conn.cursor()
        cur.execute("UPDATE sales SET status='cancelled' WHERE id=%s", (sale_id,))
        cur.close()

    report = ctx["reports"].sales_report(TODAY, TODAY, ctx["warehouse"])
    assert report.totals["invoices"] == 1
    assert report.totals["billed"] == Decimal("1260.00")
    assert ctx["reports"].profit_loss(TODAY, TODAY, ctx["warehouse"])["cogs"] \
        == Decimal("900.00")


# ------------------------------------------------------------ other reports

def test_top_products_counts_sales_net_of_returns(ctx):
    sale_id = _sell(ctx, rice=10, oil=3)
    item_id = next(l.sale_item_id for l in ctx["returns"].returnable_quantities(sale_id)
                   if l.product_id == ctx["rice"])
    ctx["returns"].create(ReturnDraft(sale_id, [ReturnLine(item_id, Decimal("2"))]))

    report = ctx["reports"].top_products(TODAY, TODAY)
    rice_row = next(r for r in report.rows if r["sku"] == f"{MARK}-RICE")
    assert Decimal(str(rice_row["sold"])) == Decimal("10.000")
    assert Decimal(str(rice_row["returned"])) == Decimal("2.000")
    assert rice_row["net_sold"] == Decimal("8.000")


def test_slow_moving_lists_stock_that_has_not_sold(ctx):
    report = ctx["reports"].slow_moving_products(days=30)
    skus = {r["sku"] for r in report.rows}
    assert f"{MARK}-OIL" in skus, "oil was bought but never sold"


def test_a_product_that_just_sold_is_not_slow_moving(ctx):
    _sell(ctx, rice=1, oil=1)
    report = ctx["reports"].slow_moving_products(days=30)
    assert f"{MARK}-OIL" not in {r["sku"] for r in report.rows}


def test_returns_report_totals(ctx):
    sale_id = _sell(ctx, rice=6)
    item_id = ctx["returns"].returnable_quantities(sale_id)[0].sale_item_id
    ctx["returns"].create(ReturnDraft(sale_id, [ReturnLine(item_id, Decimal("2"))]))
    report = ctx["reports"].sale_return_report(TODAY, TODAY)
    mine = [r for r in report.rows if r["customer"] and "Walk-in" in r["customer"]]
    assert sum(Decimal(str(r["total_amount"])) for r in mine) >= Decimal("840.00")


def test_an_empty_period_reports_zero_not_an_error(ctx):
    long_ago = TODAY - timedelta(days=400)
    report = ctx["reports"].sales_report(long_ago, long_ago)
    assert report.rows == []
    assert report.totals["billed"] == Decimal("0.00")
    assert ctx["reports"].profit_loss(long_ago, long_ago)["gross_profit"] == Decimal("0.00")


def test_a_backwards_date_range_is_refused(ctx):
    from smartpos.core.exceptions import ValidationError
    with pytest.raises(ValidationError, match="after the end date"):
        ctx["reports"].run("sales", TODAY, TODAY - timedelta(days=1))


def test_every_listed_report_runs(ctx):
    _sell(ctx, rice=2, oil=1)
    for key, label in ReportService.AVAILABLE:
        report = ctx["reports"].run(key, TODAY - timedelta(days=7), TODAY)
        assert report.columns, f"{key} produced no columns"
        assert report.title, f"{key} has no title"


# -------------------------------------------------------------- dashboard

def test_dashboard_summary(ctx):
    _sell(ctx, rice=3)
    _sell(ctx, rice=2, tendered="0")             # on credit

    summary = ctx["reports"].dashboard_summary(TODAY, ctx["warehouse"])
    assert summary["invoices_today"] == 2
    assert summary["billed_today"] == Decimal("2100.00")
    assert summary["collected_today"] >= Decimal("1260.00")
    assert len(summary["trend"]) == 14
    assert summary["trend"][-1][0] == TODAY
    assert summary["stock_value"] > Decimal("0")


# ---------------------------------------------------------------- exports

def test_the_exports_carry_the_same_rows_as_the_report(ctx, tmp_path):
    """Acceptance criterion."""
    from openpyxl import load_workbook
    from pypdf import PdfReader

    _sell(ctx, rice=3)
    _sell(ctx, rice=1, oil=2)
    report = ctx["reports"].sales_report(TODAY, TODAY, ctx["warehouse"])
    assert report.rows, "the test needs rows to compare"

    pdf_path = exporters.to_pdf(report.title, report.columns, report.rows,
                                tmp_path / "report.pdf", report.totals, report.note)
    xlsx_path = exporters.to_excel(report.title, report.columns, report.rows,
                                   tmp_path / "report.xlsx", report.totals, report.note)

    text = " ".join(page.extract_text() for page in PdfReader(str(pdf_path)).pages)
    text = text.replace("\n", " ")
    for row in report.rows:
        assert row["invoice_no"] in text, "an invoice is missing from the PDF"

    sheet = load_workbook(xlsx_path).active
    cells = {str(c.value) for row in sheet.iter_rows() for c in row if c.value is not None}
    for row in report.rows:
        assert row["invoice_no"] in cells, "an invoice is missing from the spreadsheet"

    # Money must arrive as numbers a spreadsheet can sum, not as text.
    header_row = next(r for r in sheet.iter_rows()
                      if any(c.value == "Grand total" for c in r))
    column = next(c.column for c in header_row if c.value == "Grand total")
    values = [sheet.cell(row=r, column=column).value
              for r in range(header_row[0].row + 1, sheet.max_row + 1)]
    numbers = [v for v in values if isinstance(v, (int, float))]
    assert numbers, "grand_total did not export as numbers"
    assert sum(numbers) == pytest.approx(float(report.totals["billed"]))


def test_an_empty_report_still_exports(ctx, tmp_path):
    long_ago = TODAY - timedelta(days=400)
    report = ctx["reports"].sales_report(long_ago, long_ago)
    assert exporters.to_pdf(report.title, report.columns, report.rows,
                            tmp_path / "empty.pdf").is_file()
    assert exporters.to_excel(report.title, report.columns, report.rows,
                              tmp_path / "empty.xlsx").is_file()


# -------------------------------------------------------------- performance

@pytest.mark.slow
def test_a_twelve_month_report_over_10000_sales_is_under_three_seconds(ctx):
    """Acceptance criterion."""
    warehouse, rice = ctx["warehouse"], ctx["rice"]
    walk_in = ctx["customers"].get_walk_in().id
    start = TODAY - timedelta(days=364)

    sales_rows, item_rows = [], []
    for n in range(10_000):
        when = datetime.combine(start + timedelta(days=n % 365), datetime.min.time())
        sales_rows.append((f"{MARK}-BULK-{n}", walk_in, warehouse, when,
                           "420.00", "420.00", "420.00"))
    with connection.transaction() as conn:
        cur = conn.cursor()
        cur.executemany(
            "INSERT INTO sales (invoice_no, customer_id, warehouse_id, sale_date,"
            " subtotal, grand_total, paid_amount, status, payment_status)"
            " VALUES (%s,%s,%s,%s,%s,%s,%s,'completed','paid')", sales_rows)
        cur.execute("SELECT id FROM sales WHERE invoice_no LIKE %s", (f"{MARK}-BULK-%",))
        ids = [r[0] for r in cur.fetchall()]
        for sale_id in ids:
            item_rows.append((sale_id, rice, "1.000", "420.00", "300.00", "420.00"))
        cur.executemany(
            "INSERT INTO sale_items (sale_id, product_id, quantity, unit_price,"
            " unit_cost, line_total) VALUES (%s,%s,%s,%s,%s,%s)", item_rows)
        cur.close()

    timings = {}
    for label, call in (
        ("sales totals", lambda: ctx["reports"].sales_report(start, TODAY, warehouse).totals),
        ("profit and loss", lambda: ctx["reports"].profit_loss(start, TODAY, warehouse)),
        ("top products", lambda: ctx["reports"].top_products(start, TODAY)),
    ):
        began = time.perf_counter()
        call()
        timings[label] = time.perf_counter() - began

    for label, seconds in timings.items():
        print(f"\n    {label:<18} {seconds * 1000:7.1f} ms over 10,000 sales")
        assert seconds < 3.0, f"{label} took {seconds:.2f}s"

    profit = ctx["reports"].profit_loss(start, TODAY, warehouse)
    assert profit["gross_profit"] >= Decimal("1200000.00"), "10,000 x 120 margin"


# ------------------------------------------------------------- integrity

def test_no_payment_row_points_at_a_document_that_is_gone(db):
    """A guard against the whole suite, not just this module.

    An orphaned payment inflates "collected today" and skews every due, and it
    is invisible until a figure looks wrong on the dashboard.
    """
    with connection.get_connection() as conn:
        cur = conn.cursor(dictionary=True)
        try:
            cur.execute(
                "SELECT id, reference_type, reference_id, amount "
                "  FROM customer_payments "
                " WHERE (reference_type = 'sale' "
                "        AND reference_id NOT IN (SELECT id FROM sales)) "
                "    OR (reference_type = 'sale_return' "
                "        AND reference_id NOT IN (SELECT id FROM sale_returns))")
            orphans = cur.fetchall()
            cur.execute(
                "SELECT id FROM supplier_payments "
                " WHERE reference_type = 'purchase' "
                "   AND reference_id NOT IN (SELECT id FROM purchases)")
            orphans += cur.fetchall()
        finally:
            cur.close()
    assert orphans == [], f"{len(orphans)} payment row(s) reference a deleted document"
