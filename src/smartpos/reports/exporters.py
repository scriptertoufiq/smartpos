"""Export report rows to PDF and Excel.

Both exporters take the same rows the screen is showing, so a spreadsheet and a
printout of the same report cannot disagree.
"""

from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

from smartpos.core.exceptions import ValidationError
from smartpos.core.logging import get_logger

log = get_logger(__name__)

NUMERIC_HINTS = ("total", "amount", "price", "cost", "value", "due", "payable",
                 "revenue", "profit", "paid", "discount", "tax", "quantity",
                 "sold", "returned", "on_hand", "stock")


def _is_numeric(column: str) -> bool:
    return any(hint in column.lower() for hint in NUMERIC_HINTS)


def _label(column: str) -> str:
    return column.replace("_", " ").strip().capitalize()


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, Decimal):
        return f"{value:,.2f}"
    if isinstance(value, datetime):
        return value.strftime("%d %b %Y, %H:%M")
    if isinstance(value, date):
        return value.strftime("%d %b %Y")
    if isinstance(value, float):
        return f"{value:,.2f}"
    return str(value)


def to_pdf(title: str, columns: list[str], rows: list[dict], output: Path,
           totals: dict[str, Any] | None = None, note: str = "") -> Path:
    """Render rows as a titled PDF table.

    Landscape when there are many columns, because a report squeezed onto
    portrait A4 is unreadable and gets printed twice.
    """
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
    )

    if not columns:
        raise ValidationError("The report has no columns to print.")

    output.parent.mkdir(parents=True, exist_ok=True)
    pagesize = landscape(A4) if len(columns) > 6 else A4
    document = SimpleDocTemplate(
        str(output), pagesize=pagesize,
        leftMargin=14 * mm, rightMargin=14 * mm,
        topMargin=14 * mm, bottomMargin=14 * mm, title=title)

    styles = getSampleStyleSheet()
    story = [Paragraph(title, styles["Title"])]
    story.append(Paragraph(
        f"Generated {datetime.now().strftime('%d %b %Y, %H:%M')} · "
        f"{len(rows)} row(s)", styles["Normal"]))
    if note:
        story.append(Paragraph(note, styles["Italic"]))
    story.append(Spacer(1, 8 * mm))

    data = [[_label(c) for c in columns]]
    for row in rows:
        data.append([_text(row.get(c)) for c in columns])

    if not rows:
        data.append(["No data for this period."] + [""] * (len(columns) - 1))

    table = Table(data, repeatRows=1)
    style = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f2937")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 7.5),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#d1d5db")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1),
         [colors.white, colors.HexColor("#f3f4f6")]),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]
    for index, column in enumerate(columns):
        if _is_numeric(column):
            style.append(("ALIGN", (index, 1), (index, -1), "RIGHT"))
    table.setStyle(TableStyle(style))
    story.append(table)

    if totals:
        story.append(Spacer(1, 6 * mm))
        summary = " · ".join(f"<b>{_label(k)}</b> {_text(v)}"
                             for k, v in totals.items())
        story.append(Paragraph(summary, styles["Normal"]))

    document.build(story)
    log.info("Report PDF written: %s (%d rows)", output, len(rows))
    return output


def to_excel(title: str, columns: list[str], rows: list[dict], output: Path,
             totals: dict[str, Any] | None = None, note: str = "") -> Path:
    """Write rows to an .xlsx worksheet.

    Money and quantities are written as real numbers, not text, so the shop can
    sum and filter them without retyping anything.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    if not columns:
        raise ValidationError("The report has no columns to export.")

    output.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    sheet = book.active
    sheet.title = title[:31].replace("/", "-").replace("·", "-").strip() or "Report"

    sheet.append([title])
    sheet["A1"].font = Font(size=14, bold=True)
    sheet.append([f"Generated {datetime.now().strftime('%d %b %Y, %H:%M')} · "
                  f"{len(rows)} row(s)"])
    if note:
        sheet.append([note])
        sheet.cell(row=sheet.max_row, column=1).font = Font(italic=True, size=9)
    sheet.append([])

    header_row = sheet.max_row + 1
    sheet.append([_label(c) for c in columns])
    fill = PatternFill("solid", fgColor="1F2937")
    for index in range(1, len(columns) + 1):
        cell = sheet.cell(row=header_row, column=index)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")

    for row in rows:
        values = []
        for column in columns:
            value = row.get(column)
            if isinstance(value, Decimal):
                values.append(float(value))          # a real number, not text
            elif isinstance(value, datetime):
                values.append(value.replace(tzinfo=None))
            else:
                values.append(value)
        sheet.append(values)

    for index, column in enumerate(columns, start=1):
        letter = get_column_letter(index)
        longest = max([len(_label(column))]
                      + [len(_text(r.get(column))) for r in rows] or [10])
        sheet.column_dimensions[letter].width = min(max(longest + 2, 10), 46)
        if _is_numeric(column):
            for row_index in range(header_row + 1, sheet.max_row + 1):
                cell = sheet.cell(row=row_index, column=index)
                cell.number_format = "#,##0.00"
                cell.alignment = Alignment(horizontal="right")

    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)
    if rows:
        sheet.auto_filter.ref = (
            f"A{header_row}:{get_column_letter(len(columns))}{sheet.max_row}")

    if totals:
        sheet.append([])
        for key, value in totals.items():
            sheet.append([_label(key),
                          float(value) if isinstance(value, Decimal) else value])
            sheet.cell(row=sheet.max_row, column=1).font = Font(bold=True)

    book.save(output)
    log.info("Report spreadsheet written: %s (%d rows)", output, len(rows))
    return output
